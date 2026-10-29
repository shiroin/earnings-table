#!/usr/bin/env python3
"""
企業マスター同期: ウォッチリスト（global_company_watch_300）に載っている各社の
企業マスター（*_company_master.xlsm / .xlsx）を読み、Webアプリ用の data/masters.json を作る。

  python3 sync_masters.py                       # config.json の設定で実行（GitHub Actions 用）
  python3 sync_masters.py --watch-csv W.csv --master-dir DIR   # 手元のファイルで試す

前提:
  - ウォッチリストとマスターは「リンクを知っている全員（閲覧者）」で共有されていること
  - マスターの数値は「百万（入力通貨）」単位（設定シートの input_currency）
"""
import argparse, csv, io, json, re, sys, time, datetime
from pathlib import Path

import requests
import openpyxl

ROOT = Path(__file__).resolve().parent
UA = {"User-Agent": "earnings-table-sync/1.0 (https://github.com/shiroin/earnings-table)"}

CUR = {"USD": "ドル", "JPY": "円", "EUR": "ユーロ", "CNY": "元", "KRW": "ウォン", "TWD": "台湾ドル",
       "HKD": "香港ドル", "DKK": "デンマーククローネ", "NOK": "ノルウェークローネ", "SEK": "スウェーデンクローナ",
       "CHF": "スイスフラン", "GBP": "ポンド", "INR": "ルピー", "CAD": "カナダドル", "AUD": "豪ドル", "SGD": "シンガポールドル"}
CUR_FROM_JA = {"ドル": "USD", "円": "JPY", "ユーロ": "EUR", "元": "CNY", "人民元": "CNY", "ウォン": "KRW", "台湾ドル": "TWD",
               "香港ドル": "HKD", "ポンド": "GBP"}


def file_id(url):
    m = re.search(r"/d/([a-zA-Z0-9_-]{20,})", url or "") or re.search(r"[?&]id=([a-zA-Z0-9_-]{20,})", url or "")
    return m.group(1) if m else None


# ---------------------------------------------------------------- ウォッチリスト
def read_watchlist(text):
    rows = list(csv.reader(io.StringIO(text)))
    hi = next(i for i, r in enumerate(rows) if any(c.strip() == "Ticker" for c in r))
    head = [c.strip() for c in rows[hi]]
    col = lambda name: next((i for i, h in enumerate(head) if h == name), None)
    idx = {k: col(k) for k in ["No.", "Ticker", "企業名", "日本名", "暫定Tier(個別記事化60)", "業種カテゴリー1", "業種カテゴリー2",
                               "本拠地", "上場(307)", "決算資料", "フォルダ", "Financials", "決算期", "決算日", "通貨", "単位"]}
    get = lambda r, k: (r[idx[k]].strip() if idx[k] is not None and idx[k] < len(r) else "")
    out = []
    for r in rows[hi + 1:]:
        if not get(r, "企業名") and not get(r, "Ticker"):
            continue
        # ウォッチリスト本体 = 番号があり、本拠地かマスターが入っている行（下の候補リストは除く）
        if not get(r, "No.").isdigit() or not (get(r, "本拠地") or file_id(get(r, "Financials"))):
            continue
        out.append({
            "no": get(r, "No."), "ticker": get(r, "Ticker"), "name": get(r, "企業名"), "ja": get(r, "日本名") or get(r, "企業名"),
            "tier": get(r, "暫定Tier(個別記事化60)"), "cat1": get(r, "業種カテゴリー1"), "cat2": get(r, "業種カテゴリー2"),
            "country": get(r, "本拠地"), "listed": get(r, "上場(307)"), "ir": get(r, "決算資料"),
            "folder": get(r, "フォルダ"), "master_id": file_id(get(r, "Financials")),
            "fy_end": get(r, "決算期"), "date": get(r, "決算日").replace("/", "-"),
            "currency_ja": get(r, "通貨"), "unit_in": get(r, "単位"),
        })
    return out


# ---------------------------------------------------------------- マスター
def table(ws):
    """'period' の行を見出しとして、その下の行を辞書のリストで返す"""
    rows = list(ws.iter_rows(values_only=True))
    hi = next((i for i, r in enumerate(rows) if r and str(r[0]).strip().lower() == "period"), None)
    if hi is None:
        return [], []
    head = [str(h).strip() if h is not None else "" for h in rows[hi]]
    cols = [h for h in head[1:] if h]
    data = []
    for r in rows[hi + 1:]:
        if not r or r[0] is None or str(r[0]).strip() == "":
            continue
        p = str(r[0]).strip()
        vals = {}
        for j, h in enumerate(head[1:], start=1):
            if not h:
                continue
            v = r[j] if j < len(r) else None
            if isinstance(v, str):
                v = v.replace(",", "").strip()
                try: v = float(v)
                except ValueError: v = None
            vals[h] = float(v) if isinstance(v, (int, float)) else None
        data.append({"period": p, **vals})
    # 値がすべて空の列は落とす
    cols = [c for c in cols if any(d.get(c) is not None for d in data)]
    return cols, data


def settings(wb):
    if "設定" not in wb.sheetnames:
        return {}, {}
    ws = wb["設定"]
    s, items = {}, {}
    for r in ws.iter_rows(values_only=True):
        if r and r[0] and str(r[0]).strip() not in ("setting", "会社マスター設定") and len(r) > 1:
            s[str(r[0]).strip()] = r[1]
        # 可変テーブル（種類 / 項目名 / 表示名 / カラー）: E〜H列
        if r and len(r) >= 8 and r[4] and r[5] and str(r[4]).strip() not in ("種類", "使い方"):
            items.setdefault(str(r[4]).strip(), {})[str(r[5]).strip()] = {"label": (str(r[6]).strip() if r[6] else str(r[5]).strip()),
                                                                          "color": (str(r[7]).strip() if r[7] else None)}
    return s, items


def scale_for(display_unit, currency):
    """表示単位 → (割る数, 単位ラベル)。入力は百万（通貨）"""
    cur = CUR.get(currency, currency or "")
    du = str(display_unit or "")
    if "兆" in du:   return 1_000_000, "兆" + cur
    if "十億" in du or "10億" in du: return 1_000, "10億" + cur
    if "億" in du:   return 100, "億" + cur
    if "百万" in du: return 1, "百万" + cur
    return 100, "億" + cur


PER = re.compile(r"FY\s*(\d{2,4})\s*[-_ ]?\s*Q\s*([1-4])", re.I)


def parse_period(p):
    m = PER.search(p or "")
    if not m:
        return None, None
    fy = int(m.group(1)); fy = fy + 2000 if fy < 100 else fy
    return fy, int(m.group(2))


def read_master(path_or_bytes):
    wb = openpyxl.load_workbook(io.BytesIO(path_or_bytes) if isinstance(path_or_bytes, bytes) else path_or_bytes,
                                data_only=True, read_only=False)
    s, items = settings(wb)
    out = {"company_name": s.get("company_name"), "input_currency": (s.get("input_currency") or "").strip() or None,
           "display_unit": s.get("display_unit"), "note": s.get("note")}

    def sheet(name):
        return table(wb[name]) if name in wb.sheetnames else ([], [])

    cols, q = sheet("会社全体_四半期")
    out["quarters"] = [{"period": d["period"], "revenue": d.get("revenue"), "operating_profit": d.get("operating_profit")} for d in q]
    seg_meta = items.get("セグメント", {})
    for key, name in (("seg_revenue", "セグメント売上高_四半期"), ("seg_profit", "セグメント利益_四半期")):
        c, d = sheet(name)
        out[key] = {"columns": c, "rows": d} if c else None
    out["segments"] = {k: v for k, v in seg_meta.items()}
    return out


# ---------------------------------------------------------------- 加工
def enrich(w, m):
    """ウォッチリスト1社 + マスター → Webアプリ用の1件"""
    cur = (m.get("input_currency") or CUR_FROM_JA.get(w.get("currency_ja"), "") or "").upper()
    div, unit = scale_for(m.get("display_unit"), cur)
    qs = m["quarters"]
    byp = {}
    for d in qs:
        fy, qn = parse_period(d["period"])
        if fy:
            byp[(fy, qn)] = d

    def yoy(a, b):
        if a is None or b is None or b <= 0:
            return None
        return round((a / b - 1) * 100, 1)

    periods = []
    for d in qs:
        fy, qn = parse_period(d["period"])
        prev = byp.get((fy - 1, qn)) if fy else None
        r, o = d.get("revenue"), d.get("operating_profit")
        periods.append({
            "period": d["period"], "label": f"Q{qn}" if qn else d["period"],
            "revenue": None if r is None else round(r / div, 2), "operating_profit": None if o is None else round(o / div, 2),
            "revenue_yoy": yoy(r, prev and prev.get("revenue")), "op_yoy": yoy(o, prev and prev.get("operating_profit")),
            "margin": None if (r in (None, 0) or o is None) else round(o / r * 100, 1),
        })
    periods = [p for p in periods if p["revenue"] is not None or p["operating_profit"] is not None]

    def seg(block):
        if not block:
            return None
        cols = block["columns"]
        rows = []
        for d in block["rows"]:
            vals = {c: (None if d.get(c) is None else round(d[c] / div, 2)) for c in cols}
            if any(v is not None for v in vals.values()):
                rows.append({"period": d["period"], **vals})
        return {"columns": cols, "rows": rows[-24:]}

    return {
        "ticker": w["ticker"], "name": w["name"], "ja": w["ja"], "tier": w["tier"], "cat1": w["cat1"], "cat2": w["cat2"],
        "country": w["country"], "date": w["date"], "fy_end": w["fy_end"], "master_id": w["master_id"],
        "currency": cur, "unit": unit, "note": m.get("note"),
        "periods": periods[-24:],
        "seg_revenue": seg(m.get("seg_revenue")), "seg_profit": seg(m.get("seg_profit")),
        "segments": m.get("segments") or {},
    }


def download(fid, session):
    r = session.get("https://drive.google.com/uc", params={"export": "download", "id": fid}, timeout=60, headers=UA)
    r.raise_for_status()
    if r.content[:2] != b"PK":  # ウイルススキャン確認ページなど
        m = re.search(r'name="confirm" value="([^"]+)"', r.text)
        if m:
            r = session.get("https://drive.usercontent.google.com/download",
                            params={"id": fid, "export": "download", "confirm": m.group(1)}, timeout=60, headers=UA)
        if r.content[:2] != b"PK":
            raise ValueError("Excelファイルを取得できませんでした（共有設定を確認）")
    return r.content


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--watch-csv", help="ウォッチリストのCSV（省略時は config.json のスプレッドシートから取得）")
    ap.add_argument("--master-dir", help="マスターのファイルを置いたフォルダ（試験用。ファイル名に Drive ID か Ticker を含める）")
    ap.add_argument("-o", "--out", default=str(ROOT / "data" / "masters.json"))
    a = ap.parse_args()
    cfg = json.loads((ROOT / "config.json").read_text(encoding="utf-8"))
    s = requests.Session()

    if a.watch_csv:
        text = Path(a.watch_csv).read_text(encoding="utf-8")
    else:
        w = cfg["watchlist"]
        r = s.get(f"https://docs.google.com/spreadsheets/d/{w['sheetId']}/export", params={"format": "csv", "gid": w["gid"]}, timeout=60, headers=UA)
        r.raise_for_status()
        text = r.content.decode("utf-8")
    watch = read_watchlist(text)
    print(f"ウォッチリスト: {len(watch)}社（マスターあり {sum(1 for x in watch if x['master_id'])}社）")

    old = {}
    outp = Path(a.out)
    if outp.exists():
        try: old = {c["ticker"] or c["name"]: c for c in json.loads(outp.read_text(encoding="utf-8"))["companies"]}
        except Exception: pass

    companies, errors = [], []
    for i, w in enumerate(watch, 1):
        key = w["ticker"] or w["name"]
        base = {"ticker": w["ticker"], "name": w["name"], "ja": w["ja"], "tier": w["tier"], "cat1": w["cat1"], "cat2": w["cat2"],
                "country": w["country"], "date": w["date"], "fy_end": w["fy_end"], "master_id": w["master_id"], "periods": []}
        if not w["master_id"]:
            companies.append(base); continue
        try:
            if a.master_dir:
                f = next((p for p in Path(a.master_dir).iterdir() if w["master_id"] in p.name or (w["ticker"] and w["ticker"] in p.name) or w["name"].split()[0] in p.name), None)
                if not f: companies.append(base); continue
                m = read_master(str(f))
            else:
                m = read_master(download(w["master_id"], s))
                time.sleep(0.3)
            companies.append(enrich(w, m))
        except Exception as e:
            errors.append(f"{w['name']}: {e}")
            prev = old.get(key)
            companies.append(prev if prev else base)  # 失敗したら前回の内容を残す
        if i % 25 == 0:
            print(f"  {i}/{len(watch)}")

    outp.parent.mkdir(parents=True, exist_ok=True)
    data = {"updated": datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=9))).strftime("%Y-%m-%d %H:%M"),
            "companies": companies, "errors": errors}
    outp.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    ok = sum(1 for c in companies if c.get("periods"))
    print(f"完了: 業績あり {ok}社 / エラー {len(errors)}件 → {outp}")
    for e in errors[:30]:
        print("  ", e)


if __name__ == "__main__":
    main()
